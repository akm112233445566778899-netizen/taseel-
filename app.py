import os
import re
from datetime import datetime
from io import BytesIO

import openpyxl
import pdfplumber
import streamlit as st

st.set_page_config(
    page_title="نظام تسهيل - تفريغ العهود 2026",
    page_icon="⚡",
    layout="centered",
)

# --- عرض الشعار واسم البرنامج "تسهيل" أعلى الواجهة ---
if os.path.exists("logo.png"):
    col_logo, col_title = st.columns([1, 3.5])
    with col_logo:
        st.image("logo.png", width=120)
    with col_title:
        st.title("⚡ نظام تسهيل")
        st.subheader("أداة تفريغ العهود وإصدار أمر الصرف")
        st.caption(
            "الجمعية الخيرية لتحفيظ القرآن ببريدة - نموذج الإدارة العامة 2026"
        )
else:
    st.title("⚡ نظام تسهيل")
    st.subheader("أداة تفريغ العهود وإصدار أمر الصرف")
    st.caption(
        "الجمعية الخيرية لتحفيظ القرآن ببريدة - نموذج الإدارة العامة 2026"
    )

st.divider()

VAT_RATE = 0.15


# ----------------------------------------------------------------------
# دوال مساعدة
# ----------------------------------------------------------------------
def convert_arabic_digits(text):
    """تحويل الأرقام العربية الشرقية (٠١٢٣٤٥٦٧٨٩) إلى أرقام إنجليزية"""
    if not text:
        return ""
    arabic_digits = "٠١٢٣٤٥٦٧٨٩"
    english_digits = "0123456789"
    return str(text).translate(str.maketrans(arabic_digits, english_digits))


def tafqeet_arabic(amount):
    """تحويل المبلغ إلى كلمات (صيغة مبسطة: أرقام + فقط لا غير)"""
    try:
        if amount <= 0:
            return ""
        amount = round(amount, 2)
        riyal = int(amount)
        halalas = int(round((amount - riyal) * 100))
        text = f"{riyal} ريال سعودي"
        if halalas > 0:
            text += f" و {halalas} هللة"
        return text + " فقط لا غير"
    except Exception:
        return ""


def get_arabic_month(month_num):
    months = {
        1: "يناير",
        2: "فبراير",
        3: "مارس",
        4: "أبريل",
        5: "مايو",
        6: "يونيو",
        7: "يوليو",
        8: "أغسطس",
        9: "سبتمبر",
        10: "أكتوبر",
        11: "نوفمبر",
        12: "ديسمبر",
    }
    return months.get(month_num, "")


def to_float(s):
    return float(convert_arabic_digits(s).replace(",", "").strip())


# ----------------------------------------------------------------------
# استخراج التاريخ (معدّل لدعم التواريخ قصيرة الخانات مثل 9/20/2026)
# ----------------------------------------------------------------------
def extract_invoice_date(full_text):
    found = set()

    # 1. البحث عن صيغة DD-MM-YYYY أو M/D/YYYY
    for d, m, y in re.findall(r"(?<!\d)(\d{1,2})[-/](\d{1,2})[-/](\d{4})(?!\d)", full_text):
        try:
            dt = datetime(int(y), int(m), int(d))
            if 2000 <= dt.year <= 2100:
                found.add(dt)
        except ValueError:
            try:
                dt = datetime(int(y), int(d), int(m))
                if 2000 <= dt.year <= 2100:
                    found.add(dt)
            except ValueError:
                pass

    # 2. البحث عن صيغة YYYY-MM-DD
    for y, m, d in re.findall(r"(?<!\d)(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?!\d)", full_text):
        try:
            dt = datetime(int(y), int(m), int(d))
            if 2000 <= dt.year <= 2100:
                found.add(dt)
        except ValueError:
            pass

    dates = sorted(found)
    if not dates:
        return ""
    return dates[-1].strftime("%Y-%m-%d")


# ----------------------------------------------------------------------
# استخراج المبلغ
# ----------------------------------------------------------------------
def extract_net_amount(full_text):
    amt_re = r"(?<![\d,.])(?:\d{1,3}(?:,\d{3})+|\d+)\.\d{2}(?!\d)"
    nums = set()
    for x in re.findall(amt_re, full_text):
        try:
            nums.add(to_float(x))
        except ValueError:
            pass

    best = 0.0
    for net in nums:
        if net < 1:
            continue
        vat = round(net * VAT_RATE, 2)
        total = round(net + vat, 2)
        has_vat = any(abs(vat - x) < 0.02 for x in nums)
        has_total = any(abs(total - x) < 0.02 for x in nums)
        if has_vat and has_total and net > best:
            best = net
    return best


# ----------------------------------------------------------------------
# تحليل الفاتورة (شامل لشركة الكهرباء والموردين الآخرين كبردية)
# ----------------------------------------------------------------------
def parse_pdf_invoice_robust(uploaded_file):
    data = {
        "inv_no": "",
        "account_no": "",
        "inv_date": "",
        "amount": 0.0,
        "tax_no": "",
        "supplier": "الشركة السعودية للكهرباء",
        "raw_text": "",
        "error": "",
    }

    if not uploaded_file.name.lower().endswith(".pdf"):
        data["error"] = "الاستخراج التلقائي يعمل مع ملفات PDF فقط، أدخل البيانات يدوياً."
        return data

    try:
        uploaded_file.seek(0)
        full_text_raw = ""
        with pdfplumber.open(uploaded_file) as pdf:
            for page in pdf.pages:
                full_text_raw += (page.extract_text() or "") + "\n"

        data["raw_text"] = full_text_raw
        full_text = convert_arabic_digits(full_text_raw)

        # 1. استخراج اسم المورد
        if "بردية" in full_text_raw or "Brdiah" in full_text_raw:
            data["supplier"] = "مؤسسة بردية"
        elif "الكهرباء" in full_text_raw:
            data["supplier"] = "الشركة السعودية للكهرباء"

        # 2. رقم الفاتورة (دعم فواتير الصيانة والمبيعات وفواتير الكهرباء)
        m_inv = re.search(r"فاتورة المبيعات\s*\n?\s*(\d+)", full_text_raw) or re.search(r"(?<!\d)(2[05]\d{10})(?!\d)", full_text)
        if m_inv:
            data["inv_no"] = m_inv.group(1)

        # 3. رقم الحساب (خاص بالكهرباء والمياه)
        m_acc = re.search(r"(?<!\d)(30\d{9})(?!\d)", full_text)
        if m_acc:
            data["account_no"] = m_acc.group(1)

        # 4. الرقم الضريبي للمورد (15 رقم يبدأ بـ 3)
        m_tax = re.search(r"(?<!\d)(3\d{14})(?!\d)", full_text)
        if m_tax:
            data["tax_no"] = m_tax.group(1)
        else:
            data["tax_no"] = "300000361310003" if "الكهرباء" in data["supplier"] else ""

        # 5. تاريخ الفاتورة
        data["inv_date"] = extract_invoice_date(full_text)

        # 6. المبلغ بدون ضريبة (يبحث عن إجمالي قبل الضريبة أولاً ثم المعادلات الحسابية)
        m_amt = re.search(r"اجمالي قبل الضريبة\s*[\n|:]*\s*(\d+\.\d{2})", full_text)
        if m_amt:
            data["amount"] = float(m_amt.group(1))
        else:
            data["amount"] = extract_net_amount(full_text)

    except Exception as e:
        data["error"] = f"تعذر قراءة الملف: {e}"

    return data


# ----------------------------------------------------------------------
# الواجهة
# ----------------------------------------------------------------------
dar_list = [
    "إدارة الدور النسائية",
    "دار حصة",
    "دار التقوى",
    "دار الصفا",
    "دار سمية بنت الخياط",
    "دار منيرة المانع",
    "دار الفتاة",
    "دار الفايزية",
    "دار الهدى",
    "دار الضياء",
    "دار زينب",
    "دار العليا",
    "دار الخنساء",
    "التاليات",
    "خولة بنت حكيم",
    "غراس",
    "دار بينات",
    "الأفق",
    "الوقار",
    "مركز محكم",
    "دار الهداية",
    "الرابية",
    "دار الطرفية",
]

st.subheader("اختر الدار")
dar_name = st.selectbox("اختر اسم الدار المعنية", dar_list)

st.subheader("1️⃣ رفع الفواتير (حتى 5 فواتير)")
uploaded_files = st.file_uploader(
    "التقط صور الفواتير أو اسحب ملفات الـ PDF هنا (حتى 5 فواتير معاً)",
    type=["pdf", "png", "jpg", "jpeg"],
    accept_multiple_files=True,
)

if uploaded_files:
    if len(uploaded_files) > 5:
        st.warning("⚠️ تم تحديد أكثر من 5 فواتير. سيتم معالجة أول 5 فواتير فقط.")
        uploaded_files = uploaded_files[:5]

    st.success(f"✅ تم استلام {len(uploaded_files)} فاتورة بنجاح!")
    st.divider()

    st.subheader("2️⃣ مراجعة وتحديد البيانات")

    invoices_data = []

    for idx, file in enumerate(uploaded_files):
        st.markdown(f"### 📄 الفاتورة رقم ({idx + 1}): `{file.name}`")

        with st.spinner(f"🔍 جاري استخراج البيانات للفاتورة #{idx + 1}..."):
            ext_data = parse_pdf_invoice_robust(file)

        if ext_data.get("error"):
            st.warning(f"⚠️ {ext_data['error']}")
        else:
            if not ext_data["inv_date"]:
                st.warning("⚠️ لم يتم العثور على تاريخ الفاتورة، أدخله يدوياً.")
            if not ext_data["amount"]:
                st.warning("⚠️ لم يتم العثور على المبلغ، أدخله يدوياً.")

        with st.expander(f"عرض النص المستخرج للفاتورة #{idx + 1} (للفحص)"):
            st.text(ext_data.get("raw_text", ""))

        col1, col2 = st.columns(2)
        with col1:
            supplier = st.text_input(
                f"اسم المورد #{idx+1}",
                value=ext_data.get("supplier", ""),
                key=f"sup_{idx}",
            )
            inv_no = st.text_input(
                f"رقم الفاتورة #{idx+1}",
                value=str(ext_data.get("inv_no", "")),
                key=f"inv_{idx}",
            )
            account_no = st.text_input(
                f"رقم الحساب #{idx+1}",
                value=str(ext_data.get("account_no", "")),
                key=f"acc_{idx}",
            )
            inv_date = st.text_input(
                f"تاريخ الفاتورة #{idx+1} (YYYY-MM-DD)",
                value=str(ext_data.get("inv_date", "")),
                key=f"date_{idx}",
            )

        with col2:
            tax_no = st.text_input(
                f"الرقم الضريبي #{idx+1}",
                value=str(ext_data.get("tax_no", "")),
                key=f"tax_{idx}",
            )
            amount = st.number_input(
                f"المبلغ بدون الضريبة #{idx+1}",
                value=float(ext_data.get("amount", 0.0) or 0.0),
                step=0.01,
                key=f"amt_{idx}",
            )

            # اختيار البند الافتراضي ذكياً بناءً على نوع الفاتورة
            default_cat_index = 3 if "بردية" in supplier else 0  # الصيانة كافتراضي لبردية، والكهرباء لغيرها

            category = st.selectbox(
                f"بند الحساب #{idx+1}",
                [
                    "الكهرباء",
                    "المياة",
                    "الهاتف والبريد",
                    "الصيانة",
                    "قرطاسيات ومطبوعات",
                    "نظافة",
                    "أخرى",
                ],
                index=default_cat_index,
                key=f"cat_{idx}",
            )

            statement = st.text_input(
                f"البيان #{idx+1}",
                value=f"سداد فاتورة {category} - {dar_name}",
                key=f"st_{idx}_{dar_name}_{category}",
            )

            estimated_total = round(amount + round(amount * VAT_RATE, 2), 2)
            written_amount = tafqeet_arabic(estimated_total)
            written_text = st.text_input(
                f"المبلغ كتابةً (تفقيط) #{idx+1}",
                value=written_amount,
                key=f"writ_{idx}_{amount}",
            )

        invoices_data.append({
            "supplier": supplier,
            "inv_no": inv_no,
            "account_no": account_no,
            "inv_date": inv_date,
            "amount": amount,
            "tax_no": tax_no,
            "statement": statement,
            "category": category,
            "written_amount": written_text,
        })
        st.markdown("---")

    st.subheader("3️⃣ المعالجة والإرسال")

    if st.button("🚀 تفريغ العهدة وتوزيع البنود إلى ملفات Excel"):
        excel_path = "نموذج عهدة الادارة العامة 2026.xlsx.xlsx"

        if os.path.exists(excel_path):
            try:
                cat_col_map = {
                    "الهاتف والبريد": "I",
                    "الكهرباء": "J",
                    "المياة": "K",
                    "الصيانة": "O",
                    "قرطاسيات ومطبوعات": "P",
                    "نظافة": "R",
                    "أخرى": "W",
                }

                # تجميع الفواتير حسب البند
                grouped_invoices = {}
                for item in invoices_data:
                    cat = item["category"]
                    if cat not in grouped_invoices:
                        grouped_invoices[cat] = []
                    grouped_invoices[cat].append(item)

                st.success(f"🎉 تم تقسيم الفواتير بنجاح إلى ({len(grouped_invoices)}) بند/ملف!")

                for cat_name, items in grouped_invoices.items():
                    wb = openpyxl.load_workbook(excel_path)
                    ws = wb.active

                    first_inv_date_str = items[0]["inv_date"].strip()
                    try:
                        parsed_date = datetime.strptime(first_inv_date_str, "%Y-%m-%d")
                    except Exception:
                        parsed_date = datetime.today()

                    inv_month_name = get_arabic_month(parsed_date.month)
                    inv_year = parsed_date.year

                    ws["A2"] = (
                        f"تفريغ عهدة : {dar_name} - بند ({cat_name}) شهر {inv_month_name} {inv_year} م -"
                    )

                    start_row = 6

                    for idx, item in enumerate(items):
                        current_row = start_row + idx

                        ws[f"C{current_row}"] = item["statement"]
                        ws[f"D{current_row}"] = item["account_no"]

                        raw_inv = str(item["inv_no"]).strip()
                        ws[f"E{current_row}"] = (
                            int(raw_inv) if raw_inv.isdigit() else raw_inv
                        )

                        try:
                            inv_dt = datetime.strptime(item["inv_date"].strip(), "%Y-%m-%d")
                            ws[f"F{current_row}"] = inv_dt
                            ws[f"F{current_row}"].number_format = "yyyy-mm-dd"
                        except Exception:
                            ws[f"F{current_row}"] = item["inv_date"]

                        ws[f"G{current_row}"] = item["tax_no"]
                        ws[f"G{current_row}"].number_format = "@"
                        ws[f"H{current_row}"] = item["supplier"]

                        col_letter = cat_col_map.get(cat_name, "J")
                        ws[f"{col_letter}{current_row}"] = float(item["amount"])

                    output_stream = BytesIO()
                    wb.save(output_stream)
                    output_stream.seek(0)

                    output_filename = f"تفريغ_عهدة_{dar_name}_{cat_name}_2026.xlsx"

                    st.download_button(
                        label=f"📥 تحميل ملف إكسل بند: [{cat_name}] ({len(items)} فاتورة)",
                        data=output_stream,
                        file_name=output_filename,
                        mime=(
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        ),
                        key=f"dl_{cat_name}",
                    )

                st.divider()
                st.link_button(
                    "🌐 الذهاب إلى منصة فنار لإنشاء أمر الصرف",
                    "https://fanar.gqb.org.sa/requests/home",
                )
            except Exception as e:
                st.error(f"حدث خطأ أثناء معالجة ملف الإكسل: {e}")
        else:
            st.error(
                f"لم يتم العثور على ملف النموذج المعتمد: `{excel_path}` في مجلد التطبيق!"
            )
